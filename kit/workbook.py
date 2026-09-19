"""The standard Sheet Artisan workbook: seven kinds of tab, always in the same order.

    Start Here | Data | <one tab per use> | Summary | Dictionary | Sources & Method | QA & Limits

    build_workbook("out/x.xlsx", spec=spec, rows=rows, start_here_md=rendered_markdown,
                   uses=[{"tab", "title", "why", "rows", "columns"?}],
                   summary=[{"title", "header", "rows", "chart"?: {"type": "bar"|"pie"|"line", "title"}}],
                   qa_rows=[(check, status, detail)], tier="public"|"private"|"none")

Scraped text is hostile until proven otherwise: strings are never written as formulas, control
characters are stripped, and cells are clipped at Excel's 32,767-character limit.
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

MAX_CELL = 32767
MAX_LINKS = 65000  # Excel's per-sheet hyperlink limit is 65,530
MAX_URL = 2000
EMAIL_LIMIT = 10 * 1024 * 1024

TITLE = Font(size=18, bold=True)
HEADING = Font(size=13, bold=True, color="1F4E79")
BOLD = Font(bold=True)
ITALIC = Font(italic=True, size=12)
LINK = Font(color="0563C1", underline="single")
STATUS_FILL = {"pass": "E2EFDA", "warn": "FFF2CC", "fail": "F8CBAD", "note": "EDEDED"}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, (dict, list)):
        v = json.dumps(v, ensure_ascii=False, default=str)
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if len(v) > MAX_CELL:
            v = v[: MAX_CELL - 1] + "…"
    return v


def _typed(v, t):
    v = _clean(v)
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if t in ("integer", "number") and isinstance(v, str):
        s = v.replace(",", "").strip()
        try:
            f = float(s)
            return int(f) if f.is_integer() else f
        except ValueError:
            return v
    if t == "bool" and isinstance(v, str):
        return {"true": True, "false": False, "yes": True, "no": False}.get(v.lower(), v)
    return v


def _put(ws, row, col, value, font=None):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # scraped "=..." stays text; never a live formula
    if font is not None:
        cell.font = font
    return cell


def _sheet_name(name, taken):
    base = re.sub(r"[\[\]:*?/\\]", "-", str(name)).strip().strip("'") or "Sheet"
    base = base[:31]
    candidate, i = base, 2
    while candidate.lower() in {t.lower() for t in taken}:
        suffix = f" ({i})"
        candidate = base[: 31 - len(suffix)] + suffix
        i += 1
    taken.append(candidate)
    return candidate


def _table_name(name, taken):
    base = re.sub(r"\W", "_", name)
    if not re.match(r"[A-Za-z_]", base):
        base = "T_" + base
    candidate, i = base[:200], 2
    while candidate in taken:
        candidate = f"{base[:190]}_{i}"
        i += 1
    taken.add(candidate)
    return candidate


def _widths(ws, headers, rows, max_width=60):
    for i, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(max_width, longest + 2))


def _write_table(ws, top, headers, rows, name, table_names, link_cols=(), style="TableStyleMedium2"):
    """Write headers + rows starting at row `top`; wrap in an Excel Table when there is data."""
    for j, h in enumerate(headers, 1):
        _put(ws, top, j, str(h), BOLD)
    links = 0
    for i, row in enumerate(rows, top + 1):
        for j, v in enumerate(row, 1):
            cell = _put(ws, i, j, v)
            if (j - 1) in link_cols and isinstance(v, str) and v.startswith(("http://", "https://")) \
                    and len(v) <= MAX_URL and links < MAX_LINKS:
                cell.hyperlink = v
                cell.font = LINK
                links += 1
    if rows:
        ref = f"A{top}:{get_column_letter(len(headers))}{top + len(rows)}"
        t = Table(displayName=_table_name(name, table_names), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
        ws.add_table(t)
    _widths(ws, headers, rows)


def _start_here(ws, md, sheet_names):
    """Render a small markdown subset; returns internal links whose tab does not exist."""
    ws.column_dimensions["A"].width = 110
    ws.sheet_view.showGridLines = False
    dangling = []
    r = 1
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            r += 1
            continue
        font = None
        if line.startswith("# "):
            line, font = line[2:], TITLE
        elif line.startswith("## "):
            line, font = line[3:], HEADING
        elif line.startswith("> "):
            line, font = line[2:], ITALIC
        elif re.match(r"^\s*[-*] ", line):
            line = "• " + re.sub(r"^\s*[-*] ", "", line)
        target = None
        m = re.search(r"\[([^\]]+)\]\(([^)]+)\)", line)
        if m:
            target = m.group(2).strip()
        line = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", line)
        whole_bold = re.fullmatch(r"\*\*(.+)\*\*", line.strip())
        line = line.replace("**", "")
        if whole_bold and font is None:
            font = BOLD
        cell = _put(ws, r, 1, _clean(line), font)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if target:
            if target.startswith("#"):
                sheet = target[1:]
                if sheet in sheet_names:
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1")
                    cell.font = LINK
                else:
                    dangling.append(sheet)
            elif target.startswith(("http://", "https://", "mailto:")):
                cell.hyperlink = target
                cell.font = LINK
        r += 1
    return dangling


def _chart(spec_chart, ws, top, n_rows, n_cols, anchor):
    kind = spec_chart.get("type", "bar")
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(kind, BarChart)()
    chart.title = spec_chart.get("title")
    data = Reference(ws, min_col=2, max_col=n_cols, min_row=top, max_row=top + n_rows)
    cats = Reference(ws, min_col=1, min_row=top + 1, max_row=top + n_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 7.5, 15
    ws.add_chart(chart, anchor)


def build_workbook(path, *, spec, rows, start_here_md, uses=(), summary=(), qa_rows=(),
                   tier="private", generated_at=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = spec.get("columns", [])
    names = [c["name"] for c in cols]
    types = {c["name"]: c.get("type", "text") for c in cols}
    url_idx = {i for i, c in enumerate(cols) if c.get("type") == "url"}
    warnings = []
    table_names = set()
    taken = ["Start Here", "Data", "Summary", "Dictionary", "Sources & Method", "QA & Limits"]
    use_names = [_sheet_name(u["tab"], taken) for u in uses]

    wb = Workbook()
    ws_start = wb.active
    ws_start.title = "Start Here"

    ws = wb.create_sheet("Data")
    data_rows = [[_typed(r.get(n), types[n]) for n in names] for r in rows]
    _write_table(ws, 1, names, data_rows, "Data", table_names, link_cols=url_idx)
    ws.freeze_panes = "A2"
    if len(rows) > MAX_LINKS:
        warnings.append(f"only the first {MAX_LINKS:,} source links are clickable (Excel limit)")

    for u, tab in zip(uses, use_names):
        wu = wb.create_sheet(tab)
        _put(wu, 1, 1, _clean(u.get("title") or tab), Font(size=14, bold=True))
        _put(wu, 2, 1, _clean(u.get("why") or ""))
        back = _put(wu, 3, 1, "← Start Here", LINK)
        back.hyperlink = Hyperlink(ref="A3", location="'Start Here'!A1")
        ucols = u.get("columns") or names
        uidx = {i for i, n in enumerate(ucols) if types.get(n) == "url"}
        urows = [[_typed(r.get(n), types.get(n, "text")) for n in ucols] for r in u.get("rows", [])]
        _write_table(wu, 5, ucols, urows, f"Use_{tab}", table_names, link_cols=uidx)
        wu.freeze_panes = "A6"

    wsum = wb.create_sheet("Summary")
    top = 1
    for s in summary:
        _put(wsum, top, 1, _clean(s["title"]), HEADING)
        header, srows = s["header"], [[_clean(v) for v in row] for row in s["rows"]]
        _write_table(wsum, top + 1, header, srows, f"Summary_{s['title']}", table_names,
                     style="TableStyleLight9")
        if s.get("chart") and srows and len(header) >= 2:
            _chart(s["chart"], wsum, top + 1, len(srows), len(header),
                   f"{get_column_letter(len(header) + 2)}{top}")
        top += max(len(srows) + 4, 17 if s.get("chart") else 0)

    wd = wb.create_sheet("Dictionary")
    drows = [[c["name"], c.get("type", "text"), c.get("definition", ""), c.get("source", ""),
              "yes" if c.get("derived") else "no",
              "AI-generated (see Sources & Method)" if c.get("ai") else ""] for c in cols]
    _write_table(wd, 1, ["Column", "Type", "Definition", "Source", "Derived", "AI"], drows,
                 "Dictionary", table_names)
    wd.column_dimensions["C"].width = 70

    wsrc = wb.create_sheet("Sources & Method")
    by_src = Counter(r.get("_source") for r in rows)
    fetched = defaultdict(list)
    for r in rows:
        if r.get("fetched_at"):
            fetched[r.get("_source")].append(str(r["fetched_at"]))
    srows = []
    for s in spec.get("sources", []):
        withheld = tier == "public" and s.get("publish") == "private"
        f = sorted(fetched.get(s["name"], []))
        span = (f[0][:10] if f[0][:10] == f[-1][:10] else f"{f[0][:10]} to {f[-1][:10]}") if f else ""
        srows.append([s["name"],
                      "withheld" if withheld else s.get("url", ""),
                      "withheld (collection method not published)" if withheld else s.get("method", ""),
                      by_src.get(s["name"], 0), span, s.get("license", ""), s.get("notes", "")])
    _write_table(wsrc, 1, ["Source", "URL", "Access method", "Rows", "Fetched", "License / terms", "Notes"],
                 srows, "Sources", table_names, link_cols={1})

    wq = wb.create_sheet("QA & Limits")
    qrows = [[_clean(c), _clean(s), _clean(d)] for c, s, d in qa_rows]
    _write_table(wq, 1, ["Check", "Status", "Detail"], qrows, "QA", table_names, style="TableStyleLight1")
    for i, (_, status, _) in enumerate(qrows, 2):
        fill = STATUS_FILL.get(str(status).lower())
        if fill:
            wq.cell(row=i, column=2).fill = PatternFill("solid", fgColor=fill)
    wq.column_dimensions["C"].width = 90

    for tab in _start_here(ws_start, start_here_md, set(wb.sheetnames)):
        warnings.append(f"Start Here links to a tab that does not exist: {tab!r} (tabs: {wb.sheetnames})")
    wb.properties.title = spec.get("title")
    wb.properties.creator = spec.get("author") or "Sheet Artisan"
    wb.properties.created = generated_at or datetime.now(timezone.utc).replace(tzinfo=None)
    wb.save(path)
    size = path.stat().st_size
    if size > EMAIL_LIMIT:
        warnings.append(f"workbook is {size / 1e6:.1f} MB, over the 10 MB email limit: attach a top-N cut")
    return {"path": str(path), "rows": len(rows), "size_bytes": size, "sheets": wb.sheetnames,
            "warnings": warnings}
