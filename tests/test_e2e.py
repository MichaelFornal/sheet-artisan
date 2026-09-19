"""A toy run from scaffold to publish plan, against a local fixture server. No network, no GitHub."""
import json
import os
import shutil
import subprocess
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import load_workbook

from tests._server import ROBOTS_ALLOW_ALL, FixtureServer
from tools import check_run, publish
from tools.new_run import scaffold

ITEMS = [{"id": i, "name": f"Widget Maker {i}", "category": ["tools", "parts", "kits"][i % 3],
          "employees": 5 + i * 7} for i in range(1, 61)]


def api_items(method, query, body, n):
    page = int(query.get("page", 1))
    chunk = ITEMS[(page - 1) * 20: page * 20]
    return 200, {}, {"items": chunk, "next": page < 3}


def item_detail(method, query, body, n):
    iid = int(query["_path"].rsplit("/", 1)[1])
    if iid == 7:
        return 404, {}, "gone"
    return 200, {}, next(i for i in ITEMS if i["id"] == iid)


def secret_list(method, query, body, n):
    return 200, {}, [{"id": 900 + i, "name": f"Hidden Maker {i}", "employees": 3} for i in range(4)]


SPEC = {
    "title": "Widget makers directory",
    "niche_slug": "widget-makers",
    "grain": "One row per widget maker",
    "key": ["source_url"],
    "sensitivity": "none",
    "tier": "public",
    "expected_rows": 60,
    "rate": 200.0,
    "spot_check_fields": ["name", "employees"],
    "limits": ["Only makers that list themselves in the directory."],
    "columns": [
        {"name": "name", "type": "text", "definition": "Maker name", "source": "example", "required": True},
        {"name": "category", "type": "text", "definition": "Directory category", "source": "example"},
        {"name": "employees", "type": "integer", "definition": "Headcount", "source": "example", "min": 0},
        {"name": "source_url", "type": "url", "definition": "Where the row came from", "source": "all",
         "required": True},
        {"name": "fetched_at", "type": "datetime", "definition": "Fetch time", "source": "all", "required": True},
    ],
    "sources": [
        {"name": "example", "url": "http://127.0.0.1/api/items", "method": "paginated JSON API",
         "license": "public directory", "publish": "public"},
        {"name": "secretsrc", "url": "http://127.0.0.1/api/secret", "method": "undocumented endpoint",
         "license": "terms silent", "publish": "private"},
    ],
}

SECRET_SOURCE = '''import os
NAME = "secretsrc"
BASE = os.environ["SA_EXAMPLE_BASE"]


def fetch(f, mode, sample_n):
    r = f.get(f"{BASE}/api/secret")
    for it in r.json():
        yield {"name": it["name"], "employees": it["employees"], "category": None,
               "source_url": f"{BASE}/secret/{it['id']}", "fetched_at": r.fetched_at}


def refetch(f, row):
    return {"name": row["name"], "employees": row["employees"]}
'''

FACTS = '''

@facts.fact("big_makers", "Makers with 200+ employees", fmt="int")
def big_makers(rows):
    return sum(1 for r in rows if (r.get("employees") or 0) >= 200)
'''

USES = '''

def uses(rows):
    big = sorted(rows, key=lambda r: -(r.get("employees") or 0))[:10]
    return [{"tab": "Biggest makers", "title": "The ten biggest makers", "why": "Start outreach here.",
             "rows": big, "columns": ["name", "employees", "source_url"]}]
'''

START = """# {{var:title}}
Built for {{var:company}} by {{var:author}}, {{var:date}}

## The headline
> {{fact:big_makers}} of the {{fact:row_count}} makers have two hundred or more staff.

## How you could use it
- [Call the biggest makers first](#Biggest makers)

## How it was built
- [The pipe that built this]({{var:repo_url}})

## Contact
{{var:contact}}
"""

EMAIL = """Subject: {{fact:big_makers}} widget makers your team has not called

Hi Pat,

I mapped {{fact:row_count}} widget makers; {{fact:big_makers}} have two hundred or more staff.
The workbook is attached; open Start Here first. The pipe is at {{var:repo_url}}.

{{var:author}}
"""


def make(pipe, target, base):
    env = {**os.environ, "SA_EXAMPLE_BASE": base}
    p = subprocess.run(["make", target], cwd=pipe, env=env, capture_output=True, text=True)
    if p.returncode != 0:
        raise AssertionError(f"make {target} failed:\n{p.stdout}\n{p.stderr}")
    return p


class DenyTerms(unittest.TestCase):
    def test_legal_suffix_stripped(self):
        from tools.new_run import deny_terms
        self.assertEqual(deny_terms("Acme Robotics, Inc.", "acme.test"),
                         ["Acme Robotics, Inc.", "Acme Robotics", "acme.test"])
        self.assertEqual(deny_terms("Toy Co", ""), ["Toy Co", "Toy"])


class EndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.runs = tempfile.mkdtemp()
        cls.server = FixtureServer({"/robots.txt": ROBOTS_ALLOW_ALL, "/api/items": api_items,
                                    "/items/*": item_detail, "/api/secret": secret_list,
                                    "/secret/*": lambda *a: (200, {}, {})})
        cls.server.__enter__()
        base = cls.server.url()
        cls.base = base
        cls.rundir = scaffold("Toy Co", domain="toyco.test", runs=cls.runs)
        pipe = cls.rundir / "pipe"
        spec = dict(SPEC, user_agent="sa-e2e/0.1")
        (pipe / "spec.json").write_text(json.dumps(spec, indent=2))
        (pipe / "src" / "sources" / "secretsrc.py").write_text(SECRET_SOURCE)
        with (pipe / "src" / "facts_def.py").open("a") as fh:
            fh.write(FACTS)
        (pipe / "src" / "uses.py").write_text(USES)
        tmpl = (pipe / "README.tmpl.md").read_text()
        tmpl = tmpl.replace("REPLACE: what decision this data serves, and why it did not exist as one table before.",
                            "Nobody had these makers in one list.")
        tmpl = tmpl.replace("REPLACE: the real engineering (pagination caps, entity matching, rate limits, messy sources).",
                            "Paging and provenance.")
        (pipe / "README.tmpl.md").write_text(tmpl)
        pkg = cls.rundir / "package"
        (pkg / "start_here.md").write_text(START)
        (pkg / "email.md").write_text(EMAIL)
        (pkg / "recipient.md").write_text("# Recipient\n- Name / role: Pat, founder\n- Route: email\n"
                                          "- Address: pat@toyco.test\n- Evidence: https://toyco.test/team\n"
                                          "- Pattern-derived: no\n")
        pj = json.loads((pkg / "package.json").read_text())
        pj.update(repo_url="https://github.com/MichaelFornal/widget-makers", contact="pat-free contact line",
                  author="Test Author")
        (pkg / "package.json").write_text(json.dumps(pj, indent=2))
        cls.sample = make(pipe, "sample", base)
        make(pipe, "full", base)
        make(pipe, "spot", base)
        make(pipe, "sheet", base)
        make(pipe, "readme", base)

    @classmethod
    def tearDownClass(cls):
        cls.server.__exit__(None, None, None)
        shutil.rmtree(cls.runs, ignore_errors=True)

    def test_scaffold_layout(self):
        for rel in ("run.json", "run.md", "package/package.json", "pipe/kit/fetch.py", "pipe/Makefile",
                    "pipe/.git", "out"):
            self.assertTrue((self.rundir / rel).exists(), rel)

    def test_sample_workbook(self):
        books = list((self.rundir / "out").glob("*_SAMPLE_*.xlsx"))
        self.assertEqual(len(books), 1)
        rows = [l for l in (self.rundir / "pipe/data/rows.sample.jsonl").read_text().splitlines() if l]
        self.assertEqual(len(rows), 25 + 4)  # SAMPLE_N from the paged source + the small private one

    def test_full_workbook(self):
        book = next(b for b in (self.rundir / "out").glob("*.xlsx") if "_SAMPLE_" not in b.name)
        self.assertTrue(book.name.startswith("ToyCo_widget-makers_"))
        wb = load_workbook(book)
        self.assertEqual(wb.sheetnames, ["Start Here", "Data", "Biggest makers", "Summary", "Dictionary",
                                         "Sources & Method", "QA & Limits"])
        self.assertEqual(wb["Data"].max_row, 64 + 1)
        start = [c.value for c in wb["Start Here"]["A"] if c.value]
        self.assertIn("Built for Toy Co by Test Author", " ".join(start))
        link = next(c for c in wb["Start Here"]["A"] if c.value == "• Call the biggest makers first")
        self.assertEqual(link.hyperlink.location, "'Biggest makers'!A1")
        srcs = {r[0].value: [c.value for c in r] for r in wb["Sources & Method"].iter_rows(min_row=2)}
        self.assertIn("withheld", srcs["secretsrc"][1])  # private method withheld even in the emailed copy
        qa = [[c.value for c in r] for r in wb["QA & Limits"].iter_rows(min_row=2)]
        self.assertIn("live spot-check", [r[0] for r in qa])

    def test_spot_check_found_the_gone_row(self):
        qa = json.loads((self.rundir / "pipe/data/qa.full.json").read_text())
        self.assertTrue(qa["passed"])
        self.assertEqual(qa["spot_check"]["sampled"], 20)

    def test_email_rendered_with_facts(self):
        email = (self.rundir / "out" / "email.md").read_text()
        self.assertIn("I mapped 64 widget makers", email)
        self.assertNotIn("{{", email)

    def test_check_passes_then_catches_a_stray_number(self):
        results = check_run.run_checks("toy-co", "public", self.runs)
        failed = [(n, d) for n, ok, d in results if not ok]
        self.assertEqual(failed, [])
        email = self.rundir / "package" / "email.md"
        original = email.read_text()
        try:
            email.write_text(original + "\nWe are 3x faster than anyone.\n")
            failed = {n for n, ok, d in check_run.run_checks("toy-co", "public", self.runs) if not ok}
            self.assertIn("every number is a computed fact", failed)
            self.assertIn("email rendered from the current draft", failed)
        finally:
            email.write_text(original)

    def test_public_stage_strips_private_source_and_sample_rows(self):
        plan = publish.stage("toy-co", "public", runs=self.runs)
        dest = plan["dest"]
        self.assertEqual(plan["stripped"], ["secretsrc"])
        self.assertFalse((dest / "src/sources/secretsrc.py").exists())
        spec = json.loads((dest / "spec.json").read_text())
        hidden = next(s for s in spec["sources"] if s["name"] == "secretsrc")
        self.assertTrue(hidden["withheld"])
        self.assertNotIn("api/secret", json.dumps(spec))
        self.assertEqual(spec["repo_url"], "https://github.com/MichaelFornal/widget-makers")
        sample = (dest / "sample/sample_rows.csv").read_text()
        self.assertNotIn("Hidden Maker", sample)
        self.assertEqual(plan["sample_rows"], 25)
        self.assertFalse((dest / "data").exists())
        self.assertEqual(plan["scan"]["fail"], [], plan["scan"])

    def test_public_scan_catches_the_target_name(self):
        readme = self.rundir / "pipe" / "README.md"
        original = readme.read_text()
        try:
            readme.write_text(original + "\nBuilt for Toy Co.\n")
            plan = publish.stage("toy-co", "public", runs=self.runs)
            self.assertTrue(any(k.startswith("deny:") for _, _, k, _ in plan["scan"]["fail"]))
            plan = publish.stage("toy-co", "private", runs=self.runs)
            self.assertEqual(plan["scan"]["fail"], [])
        finally:
            readme.write_text(original)

    def test_sensitive_data_is_never_public(self):
        spec_p = self.rundir / "pipe" / "spec.json"
        original = spec_p.read_text()
        try:
            spec_p.write_text(original.replace('"sensitivity": "none"', '"sensitivity": "personal"'))
            with self.assertRaises(SystemExit) as cm:
                publish.stage("toy-co", "public", runs=self.runs)
            self.assertIn("refused", str(cm.exception))
            publish.stage("toy-co", "private", runs=self.runs)  # allowed
        finally:
            spec_p.write_text(original)

    def test_go_runs_git_and_gh_with_the_right_arguments(self):
        plan = publish.stage("toy-co", "public", runs=self.runs)
        calls = []

        class Done:
            returncode = 1  # `gh repo view` says: does not exist yet

        def fake(cmd, **kw):
            calls.append(cmd)
            return Done()
        url = publish.go(plan, runner=fake)
        self.assertEqual(url, "https://github.com/MichaelFornal/widget-makers")
        self.assertEqual(calls[0], ["gh", "repo", "view", "MichaelFornal/widget-makers"])
        self.assertIn(["gh", "repo", "create", "MichaelFornal/widget-makers", "--public", "--source", ".",
                       "--push", "--description", "Widget makers directory"], calls)

    def test_none_tier_zip(self):
        plan = publish.stage("toy-co", "none", runs=self.runs)
        z = publish.go(plan, make_zip=True, runner=lambda *a, **k: self.fail("no GitHub for tier none"))
        self.assertTrue(Path(z).exists())


if __name__ == "__main__":
    unittest.main()
