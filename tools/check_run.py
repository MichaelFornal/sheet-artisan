"""The B5 gate: is this run's package safe and true enough to send?

    ./sa check <slug> [--tier public|private|none]

Every check prints PASS/FAIL with the reason. Exit 1 on any FAIL. Nothing is sent or published.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from tools.claims_check import check as claims  # noqa: E402
from tools.common import ALLOWED_TIERS, TIERS, config, load_run  # noqa: E402

MARKER = "REPL" "ACE"
PENDING_LINK = "(link added at ship)"
EMAIL_LIMIT = 10 * 1024 * 1024


def repo_url_for(spec, name=None):
    return f"https://github.com/{config()['github_owner']}/{name or spec.get('niche_slug')}"


def run_checks(slug, tier=None, runs=None):
    path, run = load_run(slug, runs)
    pipe, pkg_dir, out = path / "pipe", path / "package", path / "out"
    results = []

    def add(name, ok, detail=""):
        results.append((name, bool(ok), detail))

    spec = json.loads((pipe / "spec.json").read_text())
    pkg = json.loads((pkg_dir / "package.json").read_text())
    tier = tier or run.get("decisions", {}).get("tier") or spec.get("tier")
    sensitivity = spec.get("sensitivity")

    # spec contract
    problems = []
    if MARKER in json.dumps(spec):
        problems.append(f"spec.json still has {MARKER} markers")
    for k in ("title", "niche_slug", "grain", "key", "columns", "sources"):
        if not spec.get(k):
            problems.append(f"spec.json missing {k}")
    names = [c.get("name") for c in spec.get("columns", [])]
    for need in ("source_url", "fetched_at"):
        if need not in names:
            problems.append(f"no {need} column")
    problems += [f"column {c.get('name')} has no definition" for c in spec.get("columns", []) if not c.get("definition")]
    for s in spec.get("sources", []):
        if s.get("publish") not in ("public", "private"):
            problems.append(f"source {s.get('name')} has no publish flag (public|private)")
        if not (pipe / "src" / "sources" / f"{s.get('name')}.py").exists():
            problems.append(f"source {s.get('name')} has no module src/sources/{s.get('name')}.py")
    add("spec.json is complete", not problems, "; ".join(problems))

    # tier vs sensitivity
    allowed = ALLOWED_TIERS.get(sensitivity)
    add("tier allowed for sensitivity", tier in TIERS and allowed and tier in allowed,
        f"tier={tier} sensitivity={sensitivity} allowed={sorted(allowed) if allowed else 'unknown sensitivity'}")

    # full crawl + QA
    rows_p, qa_p = pipe / "data" / "rows.full.jsonl", pipe / "data" / "qa.full.json"
    add("full crawl exists", rows_p.exists(), "run `make full`" if not rows_p.exists() else "")
    if qa_p.exists():
        qa = json.loads(qa_p.read_text())
        fails = [c["name"] for c in qa["checks"] if c["status"] == "fail"]
        add("QA passed", qa["passed"], "failed: " + ", ".join(fails) if fails else f"{qa['row_count']:,} rows")
        spot = qa.get("spot_check")
        add("live spot-check ran and held", spot and spot["status"] != "fail",
            "run `make spot`" if not spot else f"{spot['match']}/{spot['sampled']} match ({spot['status']})")
    else:
        add("QA passed", False, "run `make sheet`")

    # claims
    facts = out / "facts.full.json"
    docs = [pkg_dir / "start_here.md", pkg_dir / "email.md", pipe / "README.tmpl.md"]
    problems = claims(docs, facts)
    add("every number is a computed fact", not problems, "; ".join(problems[:5]))

    # unfinished drafts
    unfinished = [str(p.relative_to(path)) for p in [*docs, pkg_dir / "recipient.md", pkg_dir / "package.json"]
                  if p.exists() and MARKER in p.read_text()]
    add("no unfinished drafts", not unfinished, ", ".join(unfinished))

    # workbook
    books = sorted((b for b in out.glob("*.xlsx") if "_SAMPLE_" not in b.name), key=lambda b: b.stat().st_mtime)
    if books:
        book = books[-1]
        size_ok = book.stat().st_size <= EMAIL_LIMIT
        inputs = [p for p in (rows_p, pkg_dir / "start_here.md", pkg_dir / "package.json") if p.exists()]
        stale = [p.name for p in inputs if p.stat().st_mtime > book.stat().st_mtime]
        add("workbook built from the current crawl and drafts", not stale,
            book.name if not stale else f"older than {', '.join(stale)}: `make sheet`")
        add("workbook fits in an email", size_ok, f"{book.stat().st_size / 1e6:.1f} MB")
        from openpyxl import load_workbook
        ws = load_workbook(book, read_only=True)["Start Here"]
        text = " ".join(str(c) for row in ws.iter_rows(values_only=True) for c in row if c)
        clean = PENDING_LINK not in text and MARKER not in text
        add("workbook has no placeholder text", clean,
            "" if clean else "Start Here still has a placeholder: finish package/ and `make sheet`")
    else:
        add("workbook built from the current crawl and drafts", False, "no full workbook in out/: `make sheet`")

    # email rendered and current
    email = out / "email.md"
    if email.exists():
        stale = any(email.stat().st_mtime < p.stat().st_mtime
                    for p in (pkg_dir / "email.md", pkg_dir / "package.json"))
        add("email rendered from the current draft", not stale, "re-run `make sheet`" if stale else "out/email.md")
    else:
        add("email rendered from the current draft", False, "run `make sheet` after the draft is finished")

    # repo link matches tier
    link = pkg.get("repo_url") or ""
    if tier in ("public", "private"):
        want = repo_url_for(spec)
        add("repo link matches the repo that will be published", link == want,
            link if link == want else f"package.json repo_url={link!r}, expected {want!r}")
    else:
        add("no repo link for a no-GitHub run", "github.com" not in link, f"repo_url={link!r}")
    rendered = [p for p in out.glob("*.md") if PENDING_LINK in p.read_text()]
    add("no placeholder links left", not rendered, ", ".join(p.name for p in rendered))

    rec = pkg_dir / "recipient.md"
    rec_ok = rec.exists() and MARKER not in rec.read_text()
    add("recipient recorded with evidence", rec_ok, "" if rec_ok else "fill package/recipient.md")
    return results


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("slug")
    ap.add_argument("--tier", choices=TIERS)
    ap.add_argument("--runs-dir")
    args = ap.parse_args()
    results = run_checks(args.slug, args.tier, args.runs_dir)
    for name, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  ({detail})" if detail else ""))
    failed = sum(1 for _, ok, _ in results if not ok)
    print(f"check: {len(results) - failed}/{len(results)} passed")
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
